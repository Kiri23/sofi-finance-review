#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""
Excel Validator for Agentic Finance Review

Validates resumen_quincenal.xlsx:
- File exists and can be opened by openpyxl
- Expected sheets exist (Resumen Quincenal, Distribucion, Detalle Gastos Variables)
- Resumen sheet has data (not empty)

Outputs JSON decision for Claude Code Stop hook:
- {"decision": "block", "reason": "..."} to block and retry
- {} to allow completion
"""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# Log file in same directory as this script
LOG_FILE = Path(__file__).parent / "excel-validator.log"

EXPECTED_SHEETS = [
    "Resumen Quincenal",
    "Distribucion",
    "Detalle Gastos Variables",
]


def log(message: str):
    """Append timestamped message to log file."""
    timestamp = datetime.now().strftime("%H:%M:%S")
    with open(LOG_FILE, "a") as f:
        f.write(f"[{timestamp}] {message}\n")


def find_excel(target: Path) -> Path | None:
    """Find resumen_quincenal.xlsx in the target path."""
    if target.is_file() and target.suffix == ".xlsx":
        return target
    if target.is_dir():
        candidate = target / "resumen_quincenal.xlsx"
        if candidate.exists():
            return candidate
        # Search recursively
        files = list(target.rglob("resumen_quincenal.xlsx"))
        if files:
            return files[0]
    return None


def validate_excel(file_path: Path) -> list[str]:
    """Validate the Excel workbook structure."""
    errors = []

    if not file_path.exists():
        return [f"Excel file not found: {file_path}"]

    # Try to open with openpyxl
    try:
        wb = load_workbook(file_path, read_only=True)
    except Exception as e:
        return [f"Failed to open Excel file: {e}"]

    log(f"  Sheets found: {wb.sheetnames}")

    # Check expected sheets
    for sheet_name in EXPECTED_SHEETS:
        if sheet_name not in wb.sheetnames:
            errors.append(f"Missing sheet: '{sheet_name}'")
            log(f"  Missing sheet: {sheet_name}")
        else:
            log(f"  Sheet found: {sheet_name}")

    if errors:
        wb.close()
        return errors

    # Check Resumen sheet is not empty
    resumen = wb["Resumen Quincenal"]
    row_count = 0
    for row in resumen.iter_rows(max_row=100):
        if any(cell.value is not None for cell in row):
            row_count += 1
    if row_count < 5:
        errors.append(f"Resumen Quincenal sheet has only {row_count} rows with data, expected at least 5")
        log(f"  Resumen too few rows: {row_count}")
    else:
        log(f"  Resumen rows with data: {row_count}")

    # Check Distribucion sheet has data
    dist = wb["Distribucion"]
    dist_rows = 0
    for row in dist.iter_rows(max_row=100):
        if any(cell.value is not None for cell in row):
            dist_rows += 1
    if dist_rows < 2:
        errors.append(f"Distribucion sheet has only {dist_rows} rows, expected at least 2 (header + data)")
        log(f"  Distribucion too few rows: {dist_rows}")
    else:
        log(f"  Distribucion rows: {dist_rows}")

    # Check Detalle sheet has data
    detalle = wb["Detalle Gastos Variables"]
    detalle_rows = 0
    for row in detalle.iter_rows(max_row=500):
        if any(cell.value is not None for cell in row):
            detalle_rows += 1
    if detalle_rows < 2:
        errors.append(f"Detalle Gastos Variables sheet has only {detalle_rows} rows, expected at least 2")
        log(f"  Detalle too few rows: {detalle_rows}")
    else:
        log(f"  Detalle rows: {detalle_rows}")

    wb.close()
    return errors


def main():
    log("=" * 50)
    log("EXCEL VALIDATOR STOP HOOK TRIGGERED")
    log(f"sys.argv: {sys.argv}")

    # Read hook input from stdin
    try:
        stdin_data = sys.stdin.read()
        log(f"stdin_data length: {len(stdin_data)}")
    except Exception:
        pass

    # Get target from command line arg
    if len(sys.argv) > 1:
        target = Path(sys.argv[1])
        log(f"Target (from arg): {target}")
    else:
        target = Path(".")
        log(f"Target (default): {target}")

    # Find Excel file
    excel_file = find_excel(target)

    if excel_file is None:
        log("  resumen_quincenal.xlsx not found — skipping validation")
        print(json.dumps({}))
        return

    log(f"  Validating: {excel_file}")
    errors = validate_excel(excel_file)

    # Output decision JSON
    if errors:
        log(f"RESULT: BLOCK ({len(errors)} errors)")
        for err in errors:
            log(f"  {err}")
        print(json.dumps({
            "decision": "block",
            "reason": "Excel validation failed:\n" + "\n".join(errors)
        }))
    else:
        log("RESULT: PASS - Excel validation successful")
        print(json.dumps({}))


if __name__ == "__main__":
    main()
